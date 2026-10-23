from pathlib import Path
from copy import copy
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

source = Path(r"C:\Users\r.dautbekov\Downloads\perimeter_variants (3).xlsx")
target = Path(__file__).with_name("perimeter_variants_working_existing_lights.xlsx")
wb = load_workbook(source)
summary, inputs, prices, variants = wb.worksheets

sheet_title = "Увязка опор"
if sheet_title in wb.sheetnames:
    del wb[sheet_title]
ws = wb.create_sheet(sheet_title, 0)
ws.sheet_view.showGridLines = False
ws.merge_cells("A1:D1")
ws["A1"] = "Рабочая корректировка: использовать существующие опоры освещения"
ws["A1"].font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
ws["A1"].fill = PatternFill("solid", fgColor="17365D")
ws["A1"].alignment = Alignment(vertical="center")
ws.row_dimensions[1].height = 28

intro = [
    ("Решение заказчика", "Новые специализированные опоры видеонаблюдения не устанавливать. Камеры и тепловизоры размещать на существующих опорах освещения или допустимых конструкциях зданий."),
    ("Основной вариант", "Вариант 5а: тепловизионное обнаружение снаружи + оптоволоконный датчик по ограждению + фиксированные камеры с аналитикой изнутри. Это ближайшая к согласованной схеме конфигурация в исходной таблице."),
    ("Уже учтено", "Во всех шести вариантах присутствуют внутренние 21 фиксированная камера, 4 купольные и 4 PTZ. Внутренний блок повторно добавлять не нужно."),
    ("Статус сметы", "Формулы и суммы вариантов в этом рабочем экземпляре сохранены как в исходнике. Стоимость опорных строк пока НЕ пересчитана: их надо разделить на исключённую закупку новых стоек и остающиеся адаптеры, шкафы, питание, связь и монтаж на существующих опорах."),
    ("Статус размещения", "Количество и координаты существующих световых опор не подтверждены имеющимся DXF. Поэтому строки с 34 монтажными точками камер и 12 точками тепловизоров пока являются расчётными, а не привязкой к опорам на плане."),
    ("Наружные точки", "Тепловизоры за внешним ограждением допустимы только там, где есть существующая разрешённая опора/конструкция, доступ к питанию и связь. Число 12 точек не подтверждено обследованием."),
    ("Модель тепловизора", "В исходной смете приведён неполный артикул DS-2TD2637 (15 мм) и аналитический шаг 300 м. До утверждения шага требуется полный суффикс и проверка дальности VCA по человеку для точной модификации."),
    ("Контроль дальности", "Если подразумевается DS-2TD2637-15/QY, даташит указывает VCA человека 105 м. При двух камерах от узла в разные стороны шаг 300 м может оставить около 90 м между расчётными пределами. Сценарий 180 м: 19 узлов / 38 камер на 3,37 км; это +14 камер, или +21,504 млн ₸ по текущей условной цене только за камеры."),
    ("Запас сенсорного кабеля", "Для примера Senstar FP1150 даташит рекомендует 20% запаса сенсорного кабеля: 4,044 км на контур 3,37 км (плюс 0,674 км). Применять только после выбора системы и подтверждения поставщиком."),
    ("Следующее проектное действие", "Получить актуальный генплан с ID/координатами опор, высотами, постоянным питанием и возможностью прокладки ВОЛС; затем перенести камеры на ближайшие подходящие опоры, проверить сектора/слепые зоны, пересчитать монтажные узлы и смету."),
]
for row, (label, body) in enumerate(intro, start=3):
    ws.cell(row, 1, label).font = Font(name="Arial", size=10, bold=True, color="17365D")
    ws.cell(row, 2, body).font = Font(name="Arial", size=10)
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
    ws.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[row].height = 45 if row not in (6, 8) else 60

header_row = 14
headers = ["Вариант", "Строки текущей сметы", "Текущая формулировка/количество", "Что перепроверить после обследования"]
for col, value in enumerate(headers, 1):
    cell = ws.cell(header_row, col, value)
    cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="4472C4")
    cell.alignment = Alignment(wrap_text=True, vertical="center")
ws.row_dimensions[header_row].height = 32

register = [
    ("А. Аналитические камеры", "Варианты!3", "34 точки для камер периметральной цепочки", "Исключить новые стойки; подтвердить, что монтажные комплекты/шкафы и связь рассчитаны для существующих опор."),
    ("Б. Тепловизоры", "Варианты!45", "12 точек под пары тепловизоров", "Подтвердить доступные световые опоры по наружной стороне и заменить закупку стоек на фактическую адаптацию/подключение."),
    ("В. Радар + PTZ", "Варианты!87", "9 точек установки", "Если вариант остаётся для сравнения — применять то же правило существующих опор, проверить нагрузку PTZ и положение опор."),
    ("Г. Оптоволокно + камеры", "Варианты!131", "34 точки для внутренних периметральных камер", "Использовать существующие опоры; уточнить кронштейны, шкафы, питание и оптику."),
    ("5а. Тепловизоры + волокно + аналитика", "Варианты!173, 177", "12 наружных узлов + 34 внутренних узла", "Основной кандидат. Не учитывать покупку новых стоек; смету подключения/монтажа пересчитать по реальным опорам."),
    ("5б. Тепловизоры + волокно + камеры без аналитики", "Варианты!219, 223", "12 наружных узлов + 34 внутренних узла", "Применить ограничение по опорам; оставить как стоимостное сравнение с 5а."),
    ("Общий внутренний блок", "Варианты!7, 49, 91, 135, 181, 227", "29 монтажных комплектов для 21 + 4 + 4 камер", "Убедиться, что это кронштейн/ввод/защита, а не новая опора; сверить с ведомостью опор и покрытием."),
]
thin = Side(style="thin", color="D9E2F3")
for r, values in enumerate(register, start=header_row + 1):
    for c, value in enumerate(values, 1):
        cell = ws.cell(r, c, value)
        cell.font = Font(name="Arial", size=9)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.border = Border(bottom=thin)
    ws.row_dimensions[r].height = 42

ws.cell(22, 1, "Технические ссылки")
ws.cell(22, 1).font = Font(name="Arial", size=10, bold=True, color="17365D")
ws.cell(22, 2, "Hikvision DS-2TD2637-15/QY datasheet")
ws.cell(22, 2).hyperlink = "https://assets.hikvision.com/prd/normal/all/doc/m000058239/DS-2TD2637-15_QY_Datasheet_en-US_20250512.pdf"
ws.cell(22, 2).style = "Hyperlink"
ws.cell(23, 2, "Senstar FiberPatrol FP1150 datasheet")
ws.cell(23, 2).hyperlink = "https://senstar.com/wp-content/uploads/FiberPatrol_FP1150_Perimeter_Datasheet_EN.pdf"
ws.cell(23, 2).style = "Hyperlink"

for col, width in {"A": 31, "B": 42, "C": 37, "D": 62}.items():
    ws.column_dimensions[col].width = width
ws.freeze_panes = "A14"
ws.auto_filter.ref = f"A{header_row}:D{header_row + len(register)}"

inventory_title = "Реестр опор"
if inventory_title in wb.sheetnames:
    del wb[inventory_title]
iv = wb.create_sheet(inventory_title, 1)
iv.sheet_view.showGridLines = False
iv_headers = [
    "ID опоры", "Участок / сторона ограждения", "Привязка к осям / координаты",
    "Тип и высота, м", "Постоянное питание 24/7", "Защита / резерв питания",
    "ВОЛС / сеть доступны", "Шкаф / PoE-коммутатор возможен",
    "Состояние / коррозия", "Несущая способность подтверждена",
    "Обзор / препятствия", "Предлагаемое устройство", "Направление / зона обзора",
    "Доступ для обслуживания", "Решение / примечание",
]
for col, value in enumerate(iv_headers, 1):
    cell = iv.cell(1, col, value)
    cell.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="4472C4")
    cell.alignment = Alignment(wrap_text=True, vertical="center")
iv.row_dimensions[1].height = 36
for row in range(2, 102):
    for col in range(1, len(iv_headers) + 1):
        cell = iv.cell(row, col)
        cell.font = Font(name="Arial", size=9)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.border = Border(bottom=thin)
    iv.row_dimensions[row].height = 24
for col, width in enumerate([13, 25, 30, 18, 22, 24, 20, 24, 22, 26, 27, 24, 28, 22, 34], 1):
    iv.column_dimensions[__import__("openpyxl").utils.get_column_letter(col)].width = width
iv.freeze_panes = "A2"
iv.auto_filter.ref = "A1:O101"

comment = (
    "Решение заказчика: новые специализированные опоры не устанавливать; использовать существующие опоры освещения. "
    "Текущее количество является концептуальным и не привязано к реестру световых опор. "
    "Перед фиксацией сметы проверить координаты/высоту/состояние опоры, постоянное питание, сетевой ввод, "
    "несущую способность и необходимость кронштейна/шкафа. Текущие суммы не пересчитаны под это решение."
)
for r in [3, 45, 87, 131, 173, 177, 219, 223]:
    variants.cell(r, 4).comment = Comment(comment, "OpenAI Codex")
for r in [7, 49, 91, 135, 181, 227]:
    variants.cell(r, 4).comment = Comment(
        "29 комплектов для существующих опор освещения (21 дорожная + 4 купольные + 4 PTZ). "
        "Уточнить фактические опоры, постоянное питание/связь и состав комплекта. Сумма исходной модели не пересчитана.",
        "OpenAI Codex",
    )

try:
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
except AttributeError:
    pass

wb.save(target)
print(target)
